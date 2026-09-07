"""
Convert notes.md to notes.docx with proper formatting
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

def add_inline_formatting(paragraph, text):
    """Handle bold, italic, code inline formatting"""
    parts = []
    last_end = 0
    combined_pattern = r'(\*\*.*?\*\*|\*.*?\*|`.+?`)'
    
    for match in re.finditer(combined_pattern, text):
        if match.start() > last_end:
            parts.append(('normal', text[last_end:match.start()]))
        
        matched_text = match.group(0)
        if matched_text.startswith('**'):
            parts.append(('bold', matched_text[2:-2]))
        elif matched_text.startswith('*'):
            parts.append(('italic', matched_text[1:-1]))
        elif matched_text.startswith('`'):
            parts.append(('code', matched_text[1:-1]))
        
        last_end = match.end()
    
    if last_end < len(text):
        parts.append(('normal', text[last_end:]))
    
    if not parts:
        parts.append(('normal', text))
    
    for style, content in parts:
        run = paragraph.add_run(content)
        if style == 'bold':
            run.bold = True
        elif style == 'italic':
            run.italic = True
        elif style == 'code':
            run.font.name = 'Consolas'
            run.font.size = Pt(9)

def add_table_to_doc(doc, table_lines):
    """Add a markdown table to the document"""
    if len(table_lines) < 2:
        return
    
    headers = [cell.strip() for cell in table_lines[0].strip('|').split('|')]
    num_cols = len(headers)
    
    data_rows = []
    for line in table_lines[2:]:
        cells = [cell.strip() for cell in line.strip('|').split('|')]
        if len(cells) == num_cols:
            data_rows.append(cells)
    
    table = doc.add_table(rows=1 + len(data_rows), cols=num_cols)
    table.style = 'Table Grid'
    
    # Add headers
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
    
    # Add data rows
    for row_idx, row_data in enumerate(data_rows):
        for col_idx, cell_text in enumerate(row_data):
            table.rows[row_idx + 1].cells[col_idx].text = cell_text
    
    doc.add_paragraph()

def convert_notes():
    """Main conversion function"""
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    with open('notes.md', 'r', encoding='utf-8') as f:
        content = f.read()
    
    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_block_content = []
    table_lines = []
    in_table = False
    
    while i < len(lines):
        line = lines[i]
        
        # Code block handling
        if line.strip().startswith('```'):
            if in_code_block:
                # End code block
                para = doc.add_paragraph('\n'.join(code_block_content))
                for run in para.runs:
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9)
                para.paragraph_format.left_indent = Inches(0.3)
                code_block_content = []
                in_code_block = False
            else:
                in_code_block = True
            i += 1
            continue
        
        if in_code_block:
            code_block_content.append(line)
            i += 1
            continue
        
        # Table handling
        if line.strip().startswith('|') and '|' in line[1:]:
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            add_table_to_doc(doc, table_lines)
            table_lines = []
            in_table = False
        
        # Headers
        if line.startswith('# '):
            heading = doc.add_heading(line[2:], level=0)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2)
        elif line.startswith('#### '):
            doc.add_heading(line[5:], level=3)
        elif line.startswith('---'):
            para = doc.add_paragraph('_' * 50)
        # Bullet points
        elif line.strip().startswith('- '):
            para = doc.add_paragraph(style='List Bullet')
            add_inline_formatting(para, line.strip()[2:])
        # Blockquotes
        elif line.strip().startswith('> '):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.5)
            run = para.add_run(line.strip()[2:])
            run.italic = True
        # Numbered lists
        elif re.match(r'^\d+\.', line.strip()):
            para = doc.add_paragraph(style='List Number')
            text = re.sub(r'^\d+\.\s*', '', line.strip())
            add_inline_formatting(para, text)
        # Regular text
        elif line.strip():
            para = doc.add_paragraph()
            add_inline_formatting(para, line)
        else:
            # Empty line
            doc.add_paragraph()
        
        i += 1
    
    # Handle remaining table
    if in_table and table_lines:
        add_table_to_doc(doc, table_lines)
    
    doc.save('notes.docx')
    print('Successfully created notes.docx!')
    print('File saved to: notes.docx')

if __name__ == '__main__':
    convert_notes()
