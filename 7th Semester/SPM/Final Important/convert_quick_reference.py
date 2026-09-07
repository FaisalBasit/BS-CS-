"""
Convert Quick Reference Markdown to Word Document
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

def add_heading(doc, text, level):
    """Add a heading"""
    heading = doc.add_heading(text, level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return heading

def add_inline_formatting(para, text):
    """Add text with inline markdown formatting"""
    pattern = r'(\*\*.*?\*\*|\*.*?\*|`.*?`|__.*?__|_.*?_)'
    parts = re.split(pattern, text)
    
    for part in parts:
        if not part:
            continue
        
        if part.startswith('**') and part.endswith('**'):
            run = para.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('__') and part.endswith('__'):
            run = para.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
            run = para.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('_') and part.endswith('_') and not part.startswith('__'):
            run = para.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = para.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
        else:
            para.add_run(part)

def parse_line(doc, line):
    """Parse markdown line and add to document"""
    line = line.rstrip()
    
    if not line.strip():
        doc.add_paragraph()
        return
    
    # Headers
    if line.startswith('#'):
        level = len(line) - len(line.lstrip('#'))
        text = line.lstrip('#').strip()
        add_heading(doc, text, level=min(level, 9))
        return
    
    # Horizontal rule
    if line.strip() in ['---', '___', '***']:
        doc.add_paragraph('_' * 80)
        return
    
    # Unordered list
    if re.match(r'^[\s]*[-*+]\s', line):
        text = re.sub(r'^[\s]*[-*+]\s', '', line)
        para = doc.add_paragraph(style='List Bullet')
        add_inline_formatting(para, text)
        return
    
    # Ordered list
    if re.match(r'^[\s]*\d+\.\s', line):
        text = re.sub(r'^[\s]*\d+\.\s', '', line)
        para = doc.add_paragraph(style='List Number')
        add_inline_formatting(para, text)
        return
    
    # Checkbox list
    if '- [ ]' in line or '- [x]' in line or '- [X]' in line:
        checked = '- [x]' in line or '- [X]' in line
        text = re.sub(r'^[\s]*-\s*\[[ xX]\]\s*', '', line)
        para = doc.add_paragraph(style='List Bullet')
        checkbox = para.add_run('☑ ' if checked else '☐ ')
        add_inline_formatting(para, text)
        return
    
    # Blockquote
    if line.strip().startswith('>'):
        text = line.strip()[1:].strip()
        para = doc.add_paragraph()
        add_inline_formatting(para, text)
        para.paragraph_format.left_indent = Inches(0.5)
        para.style = 'Intense Quote'
        return
    
    # Regular paragraph
    para = doc.add_paragraph()
    add_inline_formatting(para, line)

def convert_to_word(md_file, docx_file):
    """Convert markdown to Word"""
    print(f"Converting {md_file} to {docx_file}...")
    
    doc = Document()
    
    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    # Read markdown
    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    for line in lines:
        parse_line(doc, line)
    
    doc.save(docx_file)
    print(f"✓ Successfully created {docx_file}")

if __name__ == "__main__":
    md_file = "quick_reference.md"
    docx_file = "quick_reference.docx"
    
    try:
        convert_to_word(md_file, docx_file)
        print("\n✓ Conversion completed successfully!")
        print(f"  Output file: {docx_file}")
    except Exception as e:
        print(f"\n✗ Error: {e}")
        import traceback
        traceback.print_exc()
